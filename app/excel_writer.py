"""用 openpyxl 生成三页 Excel 报表。

要点:
* 序号/版本号/作者/日期/操作/扩展名/统计列居中，说明与路径左对齐换行。
* 提交说明中的「更改原因/问题单号/更改措施/影响范围/桌面验证」等字段加粗并各自另起一行。
* 先写临时文件再替换目标文件，避免写一半损坏原文件。
"""

from __future__ import annotations

import math
import os
import re
from datetime import datetime
from typing import Callable, Sequence

from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .models import LogEntry

FONT_NAME = "微软雅黑"

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(name=FONT_NAME, size=10, bold=True, color="FFFFFF")
BODY_FONT = Font(name=FONT_NAME, size=10)
BOLD_FONT = Font(name=FONT_NAME, size=10, bold=True)
TITLE_FONT = Font(name=FONT_NAME, size=14, bold=True, color="1F4E78")
ALT_FILL = PatternFill("solid", fgColor="F2F7FB")
LABEL_FILL = PatternFill("solid", fgColor="DDEBF7")

THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

CENTER = Alignment(horizontal="center", vertical="center", wrap_text=False)
LEFT_TOP = Alignment(horizontal="left", vertical="top")
WRAP_TOP = Alignment(horizontal="left", vertical="top", wrap_text=True)

DETAIL_HEADERS = ["序号", "版本号", "作者", "提交日期", "提交说明", "操作", "文件路径", "扩展名"]
DETAIL_WIDTHS = [6, 9, 15, 18, 42, 8, 66, 8]
DETAIL_CENTER = {1, 2, 3, 4, 6, 8}

SUMMARY_HEADERS = ["序号", "版本号", "作者", "提交日期", "提交说明", "文件数", "新增", "修改", "删除", "文件列表"]
SUMMARY_WIDTHS = [6, 9, 15, 18, 42, 8, 7, 7, 7, 78]
SUMMARY_CENTER = {1, 2, 3, 4, 6, 7, 8, 9}

MESSAGE_COL = 5
MESSAGE_WIDTH = DETAIL_WIDTHS[MESSAGE_COL - 1]

# 提交说明中需要加粗并独立成行的字段（可按需扩充）
SECTION_KEYWORDS = (
    "更改原因",
    "修改原因",
    "问题单号",
    "缺陷单号",
    "更改措施",
    "修改措施",
    "影响范围",
    "桌面验证",
    "验证情况",
    "测试说明",
    "备注",
)
_KEY_ALT = "|".join(SECTION_KEYWORDS)
_LINE_PATTERN = re.compile(rf"^({_KEY_ALT})\s*[：:]\s*(.*)$")
_BLANK_LINES = re.compile(r"\n{3,}")


# --------------------------------------------------------------------------- #
# 提交说明分段
# --------------------------------------------------------------------------- #

def split_message_segments(message: str) -> list[tuple[str, str]]:
    """把提交说明拆成 [(加粗字段名, 内容)]，加粗字段名为空表示普通行。

    ``更改原因：xxx 问题单号：123`` 这种写在一行里的也会被拆开。
    """
    text = (message or "").replace("\r\n", "\n").replace("\r", "\n").strip()
    if not text:
        return []

    for keyword in SECTION_KEYWORDS:
        # 字段名不在行首时，先补一个换行；不处理已在行首的情况
        text = re.sub(rf"(?<!\n)[ \t]*(?={keyword}\s*[：:])", "\n", text)
    text = _BLANK_LINES.sub("\n\n", text)

    segments: list[tuple[str, str]] = []
    for raw in text.split("\n"):
        line = raw.strip()
        if not line:
            continue
        match = _LINE_PATTERN.match(line)
        if match:
            segments.append((match.group(1) + "：", match.group(2).strip()))
        else:
            segments.append(("", line))
    return segments


def _message_value(message: str):
    """生成富文本单元格值：字段名加粗，每段另起一行。"""
    segments = split_message_segments(message)
    if not segments:
        return ""
    if len(segments) == 1 and not segments[0][0]:
        return segments[0][1]

    blocks: list = []
    for index, (prefix, body) in enumerate(segments):
        if index:
            blocks.append("\n")
        if prefix:
            blocks.append(TextBlock(InlineFont(rFont=FONT_NAME, sz=10, b=True), prefix))
            if body:
                blocks.append(TextBlock(InlineFont(rFont=FONT_NAME, sz=10), body))
        elif body:
            blocks.append(TextBlock(InlineFont(rFont=FONT_NAME, sz=10), body))
    return CellRichText(blocks)


# --------------------------------------------------------------------------- #
# 通用样式
# --------------------------------------------------------------------------- #

def _set_widths(ws: Worksheet, widths: Sequence[int]) -> None:
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width


def _write_header(ws: Worksheet, headers: Sequence[str]) -> None:
    for column, title in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=column, value=title)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = BORDER
    ws.row_dimensions[1].height = 22
    ws.freeze_panes = "A2"


def _text_units(text: str) -> int:
    """中文按 2 个字符宽度估算。"""
    return sum(2 if ord(ch) > 0x2E7F else 1 for ch in text)


def _line_count(text: str, width: int) -> int:
    if not text:
        return 1
    per_line = max(8, width - 2)
    total = 0
    for part in text.replace("\r\n", "\n").split("\n"):
        total += max(1, math.ceil(_text_units(part) / per_line))
    return total


def _row_height(*pairs: tuple[str, int]) -> float:
    lines = max((_line_count(text, width) for text, width in pairs), default=1)
    return min(max(16.0, lines * 13.5 + 5), 220.0)


def _write_row(
    ws: Worksheet,
    row: int,
    values: Sequence[object],
    shade: bool,
    center_columns: set[int],
) -> None:
    for column, value in enumerate(values, start=1):
        cell = ws.cell(row=row, column=column)
        if column == MESSAGE_COL and isinstance(value, str):
            cell.value = _message_value(value)
        else:
            cell.value = value
        cell.font = BODY_FONT
        cell.border = BORDER

        if column == 4:
            cell.number_format = "yyyy-mm-dd hh:mm"
            cell.alignment = CENTER
        elif column in center_columns:
            cell.alignment = CENTER
        elif column in (5, 7, 10):
            cell.alignment = WRAP_TOP
        else:
            cell.alignment = LEFT_TOP

        if shade:
            cell.fill = ALT_FILL


# --------------------------------------------------------------------------- #
# 三个工作表
# --------------------------------------------------------------------------- #

def _build_basic_sheet(ws: Worksheet, info_items: Sequence[tuple[str, str]]) -> None:
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 96

    ws.merge_cells("A1:B1")
    title = ws["A1"]
    title.value = "SVN 日志导出报表 - 基本信息"
    title.font = TITLE_FONT
    title.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 30

    row = 3
    for label, value in info_items:
        label_cell = ws.cell(row=row, column=1, value=label)
        label_cell.font = BOLD_FONT
        label_cell.fill = LABEL_FILL
        label_cell.alignment = Alignment(horizontal="center", vertical="center")
        label_cell.border = BORDER

        value_cell = ws.cell(row=row, column=2, value=value)
        value_cell.font = BODY_FONT
        value_cell.alignment = WRAP_TOP
        value_cell.border = BORDER
        row += 1

    note_row = row + 1
    ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=2)
    note = ws.cell(row=note_row, column=1)
    note.value = (
        "说明：本报表仅统计扩展名过滤后的文件；被删除(D)的文件同样保留，便于追溯。"
        "「日志明细」按文件一行，「版本汇总」按版本一行。"
    )
    note.font = Font(name=FONT_NAME, size=9, italic=True, color="808080")
    note.alignment = WRAP_TOP

    ws.sheet_view.showGridLines = False


def _build_detail_sheet(ws: Worksheet, entries: Sequence[LogEntry]) -> None:
    _write_header(ws, DETAIL_HEADERS)
    _set_widths(ws, DETAIL_WIDTHS)

    row = 2
    index = 0
    for group, entry in enumerate(entries):
        shade = group % 2 == 1
        message = entry.summary_message()
        if not entry.files:
            index += 1
            values = [index, entry.revision, entry.author, entry.date, message, "", "", ""]
            _write_row(ws, row, values, shade, DETAIL_CENTER)
            ws.row_dimensions[row].height = _row_height((message, MESSAGE_WIDTH))
            row += 1
            continue

        for change in entry.files:
            index += 1
            values = [
                index,
                entry.revision,
                entry.author,
                entry.date,
                message,
                change.action_cn,
                change.path,
                change.ext,
            ]
            _write_row(ws, row, values, shade, DETAIL_CENTER)
            ws.row_dimensions[row].height = _row_height(
                (message, MESSAGE_WIDTH),
                (change.path, DETAIL_WIDTHS[6]),
            )
            row += 1

    last_row = max(row - 1, 1)
    ws.auto_filter.ref = f"A1:H{last_row}"


def _build_summary_sheet(ws: Worksheet, entries: Sequence[LogEntry]) -> None:
    _write_header(ws, SUMMARY_HEADERS)
    _set_widths(ws, SUMMARY_WIDTHS)

    row = 2
    for index, entry in enumerate(entries, start=1):
        file_list = "\n".join(f"{f.action_cn} {f.path}" for f in entry.files) or "(无匹配文件)"
        message = entry.summary_message()
        values = [
            index,
            entry.revision,
            entry.author,
            entry.date,
            message,
            entry.file_count,
            entry.count_action("A"),
            entry.count_action("M"),
            entry.count_action("D"),
            file_list,
        ]
        _write_row(ws, row, values, index % 2 == 0, SUMMARY_CENTER)
        ws.row_dimensions[row].height = _row_height(
            (message, MESSAGE_WIDTH),
            (file_list, SUMMARY_WIDTHS[9]),
        )
        row += 1

    last_row = max(row - 1, 1)
    ws.auto_filter.ref = f"A1:J{last_row}"

    total_row = row + 1
    ws.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=5)
    title_cell = ws.cell(row=total_row, column=1, value="合计")
    title_cell.font = BOLD_FONT
    title_cell.alignment = CENTER
    title_cell.fill = LABEL_FILL
    for column, value in (
        (6, sum(e.file_count for e in entries)),
        (7, sum(e.count_action("A") for e in entries)),
        (8, sum(e.count_action("M") for e in entries)),
        (9, sum(e.count_action("D") for e in entries)),
    ):
        cell = ws.cell(row=total_row, column=column, value=value)
        cell.font = BOLD_FONT
        cell.alignment = CENTER
        cell.fill = LABEL_FILL
    for column in range(1, 11):
        ws.cell(row=total_row, column=column).border = BORDER
    ws.row_dimensions[total_row].height = 20


# --------------------------------------------------------------------------- #
# 入口
# --------------------------------------------------------------------------- #

def write_report(
    path: str,
    info_items: Sequence[tuple[str, str]],
    entries: Sequence[LogEntry],
    on_progress: Callable[[int], None] | None = None,
) -> None:
    """生成三页报表并保存（先写临时文件再替换，失败不会破坏原文件）。"""
    workbook = Workbook()

    basic = workbook.active
    basic.title = "基本信息"
    _build_basic_sheet(basic, info_items)
    if on_progress:
        on_progress(20)

    detail = workbook.create_sheet("日志明细")
    _build_detail_sheet(detail, entries)
    if on_progress:
        on_progress(65)

    summary = workbook.create_sheet("版本汇总")
    _build_summary_sheet(summary, entries)
    if on_progress:
        on_progress(90)

    workbook.properties.title = "SVN 日志导出报表"
    workbook.properties.creator = "SVN 日志导出工具"
    workbook.properties.created = datetime.now()

    temp_path = f"{path}.tmp"
    try:
        workbook.save(temp_path)
        os.replace(temp_path, path)
    finally:
        if os.path.exists(temp_path):
            try:
                os.remove(temp_path)
            except OSError:
                pass

    if on_progress:
        on_progress(100)
