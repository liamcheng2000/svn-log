"""核心逻辑自测：python -m unittest discover -s tests -v"""

from __future__ import annotations

import os
import sys
import tempfile
import unittest
from datetime import datetime

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from openpyxl import load_workbook  # noqa: E402

from app import svn_client  # noqa: E402
from app.excel_writer import write_report  # noqa: E402
from app.models import LogEntry  # noqa: E402

SAMPLE_XML = """<?xml version="1.0" encoding="UTF-8"?>
<log>
<logentry revision="12">
<author>zhangsan</author>
<date>2024-06-30T02:12:33.123456Z</date>
<paths>
<path action="M" kind="file">/trunk/src/main.cpp</path>
<path action="A" kind="file">/trunk/src/util.H</path>
<path action="M" kind="dir">/trunk/src</path>
<path action="D" kind="file">/trunk/readme.md</path>
</paths>
<msg>修复崩溃 &amp; 优化</msg>
</logentry>
<logentry revision="11">
<author>lisi</author>
<date>2024-06-28T09:00:00.000000Z</date>
<paths>
<path action="M" kind="file">/trunk/doc/guide.txt</path>
</paths>
<msg>只改文档</msg>
</logentry>
</log>""".encode("utf-8")


class PatternTests(unittest.TestCase):
    def test_parse_patterns(self) -> None:
        self.assertEqual(svn_client.parse_patterns(".cpp;.h"), [".cpp", ".h"])
        self.assertEqual(svn_client.parse_patterns("*.cpp, H"), [".cpp", ".h"])
        self.assertEqual(svn_client.parse_patterns("cpp; h;.h"), [".cpp", ".h"])
        self.assertEqual(svn_client.parse_patterns(""), [])
        self.assertEqual(svn_client.parse_patterns("   "), [])

    def test_rev_spec(self) -> None:
        self.assertIsNone(svn_client.build_rev_spec("rev", "", ""))
        self.assertEqual(svn_client.build_rev_spec("rev", "", ""), None)
        self.assertEqual(svn_client.build_rev_spec("rev", "10", ""), "10:HEAD")
        self.assertEqual(svn_client.build_rev_spec("rev", "", "20"), "1:20")
        self.assertEqual(
            svn_client.build_rev_spec("date", "2024-01-01", "2024-06-30"),
            "{2024-01-01T00:00:00}:{2024-06-30T23:59:59}",
        )
        self.assertEqual(svn_client.build_rev_spec("date", "", "2024-06-30"), "1:{2024-06-30T23:59:59}")
        self.assertIsNone(svn_client.build_rev_spec("date", "", ""))


class ParseTests(unittest.TestCase):
    def test_parse_and_filter(self) -> None:
        entries = svn_client.parse_log_xml(SAMPLE_XML)
        self.assertEqual(len(entries), 2)
        self.assertEqual(entries[0].revision, 12)
        self.assertEqual(entries[0].author, "zhangsan")
        self.assertEqual(len(entries[0].files), 3)  # 目录被忽略
        self.assertEqual(entries[0].files[0].ext, ".cpp")
        self.assertEqual(entries[0].files[1].action_cn, "新增")
        self.assertIn("修复崩溃", entries[0].message)

        filtered = svn_client.filter_entries(entries, [".cpp", ".h"])
        self.assertEqual(len(filtered), 1)
        self.assertEqual(filtered[0].revision, 12)
        self.assertEqual(filtered[0].file_count, 2)

        kept = svn_client.filter_entries(entries, [".cpp"], keep_empty=True)
        self.assertEqual([e.revision for e in kept], [12, 11])
        self.assertEqual(kept[1].file_count, 0)


class ExcelTests(unittest.TestCase):
    def test_write_report(self) -> None:
        entries = svn_client.filter_entries(svn_client.parse_log_xml(SAMPLE_XML), [".cpp", ".h"])
        info = [
            ("生成时间", "2024-07-01 10:00:00"),
            ("SVN 地址", "https://svn.example.com/repo/trunk"),
            ("查询模式", "按版本范围"),
            ("导出文件数", "2"),
        ]
        with tempfile.TemporaryDirectory() as folder:
            path = os.path.join(folder, "out.xlsx")
            write_report(path, info, entries)
            self.assertTrue(os.path.isfile(path))

            workbook = load_workbook(path)
            self.assertEqual(workbook.sheetnames, ["基本信息", "日志明细", "版本汇总"])
            basic, detail, summary = (workbook[name] for name in workbook.sheetnames)
            self.assertIn("基本信息", str(basic["A1"].value))
            self.assertEqual(detail["A1"].value, "序号")
            self.assertEqual(detail["B2"].value, 12)
            self.assertEqual(detail["G2"].value, "/trunk/src/main.cpp")
            self.assertEqual(detail.freeze_panes, "A2")
            self.assertEqual(summary["F2"].value, 2)  # 文件数
            self.assertEqual(summary["A2"].value, 1)

    def test_empty_message_entry(self) -> None:
        entry = LogEntry(revision=1, author="a", date=datetime.now(), message="", files=[])
        with tempfile.TemporaryDirectory() as folder:
            path = os.path.join(folder, "empty.xlsx")
            write_report(path, [("x", "y")], [entry])
            self.assertTrue(os.path.isfile(path))


if __name__ == "__main__":
    unittest.main()
